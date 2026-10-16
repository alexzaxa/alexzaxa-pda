"""One-off: builds the English CV docx, matching the existing Greek CV's formatting
(same section structure, same bold/size conventions). Run once; the output is a
real file to be committed, not regenerated automatically."""
from docx import Document
from docx.shared import Pt, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Match the original Greek CV's compact layout (that file was built in Word
# with tight margins/spacing; python-docx's bare-Document defaults are much
# looser and pushed this to 2 pages).
section = doc.sections[0]
section.top_margin = Emu(502920)
section.bottom_margin = Emu(502920)
section.left_margin = Emu(594360)
section.right_margin = Emu(594360)

normal = doc.styles['Normal']
normal.font.name = 'Aptos'
normal.font.size = Emu(120650)

def _tight(p):
    p.paragraph_format.space_after = Emu(25400)
    return p

def centered_bold(text, size):
    p = _tight(doc.add_paragraph())
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    return p

def centered_plain(text):
    p = _tight(doc.add_paragraph())
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)
    return p

def heading(text):
    p = _tight(doc.add_paragraph())
    p.paragraph_format.space_before = Emu(88900)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(11)
    return p

def labeled_line(label, rest):
    p = _tight(doc.add_paragraph())
    r1 = p.add_run(label)
    r1.bold = True
    p.add_run(rest)
    return p

def plain(text):
    _tight(doc.add_paragraph(text))

centered_bold("ALEXANDROS ZACHAROPOULOS", 21)
centered_bold("IT Support • Python & Web Development • Networking", 11)
centered_plain("Chalkida, Greece  |  +30 6981868758  |  alexzaxa70@gmail.com")

heading("PROFESSIONAL PROFILE")
plain(
    "Early-career IT professional with hands-on experience in IT support, networking, "
    "Windows/Linux systems, and application development. Particularly interested in "
    "Python, web development, automation, and network infrastructure, and have built "
    "personal projects that combine software, databases, networking, and automation. "
    "Analytical thinker, quick to learn, and comfortable adapting to different technical "
    "needs. Open to both software development roles and IT support, systems, or "
    "networking positions."
)

heading("TECHNICAL SKILLS")
labeled_line("Programming & Web: ", "Python, HTML, CSS, JavaScript, Flask, REST APIs, responsive web applications")
labeled_line("Networking: ", "LAN design, TCP/IP, NAT, port forwarding, DNS, VM networking, network security fundamentals")
labeled_line("Systems: ", "Windows 11, Linux, system administration, user management, Active Directory, virtualization")
labeled_line("Development & Tools: ", "Git/GitHub, VS Code, Visual Studio, API integration, databases, automation")
labeled_line("Cybersecurity: ", "Kali Linux, penetration-testing fundamentals, vulnerability assessment, security awareness")
labeled_line("AI & Automation: ", "Python automation, data processing, AI/ML fundamentals, workflow automation")

heading("PROFESSIONAL EXPERIENCE")
labeled_line("IT / Technical Support — Halcor", "  |  Internship / Work experience")
plain("Supported day-to-day IT issues and provided technical troubleshooting.")
plain("Gained familiarity with computer systems, networks, and IT environment processes.")
plain("Built practical experience resolving technical problems and supporting end users.")
labeled_line("Sales / Customer Service", "  |  Professional experience")
plain("Customer service, communication, and understanding customer needs.")
plain("Reliability, consistency, and adaptability in a fast-paced professional environment.")

heading("PERSONAL & TECHNICAL PROJECTS")
labeled_line(
    "AlexZaxa Solutions: ",
    "Design and development of professional websites for local businesses, with a focus "
    "on responsive design, a mobile-first experience, and a modern web stack.",
)
labeled_line(
    "Restaurant PDA / Order Management: ",
    "Built an order-taking application with waiter/PDA functionality, order management, "
    "a delivery workflow, kitchen printing, and data management.",
)
labeled_line(
    "Home Lab & Infrastructure: ",
    "Hands-on work with Windows Server, Active Directory, Linux servers, virtualization, "
    "Docker, monitoring, DNS, and secure remote access.",
)

heading("EDUCATION & TRAINING")
plain("Oinofyta / Schimatari Vocational High School (EPAL) — Informatics")
plain("Certification / training in Networking")
plain("Microsoft Applied Skills — Active Directory, Azure, and GitHub Copilot")
plain("Seminars / hands-on practice in Kali Linux, CTF challenges, and cybersecurity fundamentals")

heading("LANGUAGES")
plain("Greek — Native language")
plain("English — C2 / Proficiency (Michigan & LRN)")

heading("CAREER OBJECTIVE")
plain(
    "Looking for an opportunity in IT, with a preference for IT Support, Network/System "
    "Administration, Web Development, or Python development. Open to taking on a range "
    "of technical responsibilities, learning new tools, and growing within a professional "
    "environment."
)

doc.save("assets/cv/alexandros-zacharopoulos-cv.docx")
print("wrote assets/cv/alexandros-zacharopoulos-cv.docx")
